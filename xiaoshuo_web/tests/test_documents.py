from io import BytesIO
import unittest

from document_loader import DocumentError, load_document


class DocumentTests(unittest.TestCase):
    def test_txt_utf8_bom_and_gb18030_preserve_all_characters(self):
        text = "第一章 开端\n林秋离开家乡。" * 20 + "\n最后一章 终局秘密"
        for encoding in ("utf-8", "utf-8-sig", "gb18030"):
            self.assertEqual(load_document(text.encode(encoding), "book.TXT"), text)

    def test_utf16_bom_is_detected_without_guessing_byte_order(self):
        text = "第一章 冬夜里，他捧着最后一本旧书走过桥头。\n" * 20
        for data in (text.encode("utf-16"), b"\xfe\xff" + text.encode("utf-16-be")):
            self.assertEqual(load_document(data, "book.txt"), text)
        with self.assertRaisesRegex(DocumentError, "编码无法识别"):
            load_document(b"\xff\xfe\x00", "broken.txt")

    def test_docx_includes_paragraphs_tables_and_ending_in_order(self):
        from docx import Document
        document = Document()
        document.add_paragraph("第一章 正文开始。" * 10)
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "表格中的重要线索"
        table.cell(0, 1).text = "证人证词"
        document.add_paragraph("最后一章 真相终于揭晓。")
        stream = BytesIO()
        document.save(stream)
        text = load_document(stream.getvalue(), "book.docx")
        self.assertLess(text.index("正文开始"), text.index("重要线索"))
        self.assertLess(text.index("重要线索"), text.index("真相终于揭晓"))

    def test_docx_expansion_limit(self):
        from docx import Document
        document = Document()
        document.add_paragraph("Novel text " * 20)
        stream = BytesIO()
        document.save(stream)
        with self.assertRaisesRegex(DocumentError, "解压"):
            load_document(stream.getvalue(), "book.docx", max_expanded_bytes=10)

    def make_pdf(self, texts):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter()
        for text in texts:
            page = writer.add_blank_page(width=612, height=792)
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(stream)
        output = BytesIO()
        writer.write(output)
        return output.getvalue()

    def test_pdf_extracts_all_pages(self):
        start = "Chapter One. Alice returned to the old town and discovered a hidden letter."
        ending = "Final chapter. The letter finally revealed the secret and Alice left town."
        text = load_document(self.make_pdf([start, ending]), "novel.pdf")
        self.assertIn(start, text)
        self.assertIn(ending, text)
        self.assertLess(text.index(start), text.index(ending))

    def test_long_pdf_exact_limit_retains_final_page_and_rejects_overflow(self):
        pages = ["A" * 160_000, "B" * 160_000 + " Final page secret."]
        expected = "\n\n".join(pages)
        data = self.make_pdf(pages)
        self.assertEqual(load_document(data, "long.pdf", max_chars=len(expected)), expected)
        with self.assertRaisesRegex(DocumentError, "未截断"):
            load_document(data, "long.pdf", max_chars=len(expected) - 1)

    def test_unreadable_pdf_page_is_not_silently_dropped(self):
        data = self.make_pdf(["This first page has enough text to be accepted as an entire document.", ""])
        with self.assertRaisesRegex(DocumentError, "第 2 页"):
            load_document(data, "novel.pdf")

    def test_encrypted_pdf_and_malformed_documents_are_rejected(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        writer.encrypt("secret")
        output = BytesIO()
        writer.write(output)
        with self.assertRaisesRegex(DocumentError, "解密"):
            load_document(output.getvalue(), "secret.pdf")
        for name in ("broken.docx", "broken.pdf", "binary.txt", "script.exe"):
            with self.assertRaises(DocumentError):
                load_document(b"\x00\x01" * 100, name)

    def test_over_limit_text_is_rejected(self):
        with self.assertRaisesRegex(DocumentError, "不会截断"):
            load_document(b"x" * 101, "novel.txt", max_chars=100)


if __name__ == "__main__":
    unittest.main()
